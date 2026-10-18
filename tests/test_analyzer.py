"""
Unit tests for analyzer.py.  Run from the project folder with:

    python -m unittest discover -s tests -v

Nothing here opens a window, and nothing touches your real history.json.
"""
import json
import os
import sys
import tempfile
import unittest
from unittest import mock

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import analyzer as a   # noqa: E402


# A resume that should trigger no tips at all
GOOD_RESUME = """Priya Sharma
priya@example.com | +91 98765 43210

Summary
Data analyst who turns messy sales data into clear decisions for managers.

Skills
Python, SQL, Excel, Power BI, Pandas, NumPy, Statistics, Data Visualization

Experience
Analyst Intern, ABC Retail (2024)
- Built weekly Power BI dashboards that cut reporting time by 30%.
- Analyzed 50,000 orders in SQL and Python to find the 5 best-selling products.
- Automated Excel reports with Pandas, saving the team 6 hours every week.
- Presented findings to 12 managers and helped change the discount strategy.

Projects
Customer churn analysis
- Cleaned 20,000 customer records and trained a model that reached 85% accuracy.
- Created charts with Matplotlib and Seaborn to explain the results to a team.

Certifications
- Completed a 40 hour SQL course and a Power BI course from Microsoft Learn in 2024.
- Finished the Kaggle Pandas and Data Visualization courses with 100% of exercises solved.

Education
B.Tech in Computer Science Engineering, 2024. Final year project on retail analytics.
"""


class FindSkillsTests(unittest.TestCase):
    def test_exact_names(self):
        self.assertEqual(a.find_skills("I know Python and SQL."), ["Python", "SQL"])

    def test_aliases_map_to_the_real_skill(self):
        found = a.find_skills("ML, sklearn, PowerBI, Postgres, K8s, JS")
        for skill in ["Machine Learning", "Scikit-learn", "Power BI",
                      "PostgreSQL", "Kubernetes", "JavaScript"]:
            self.assertIn(skill, found)

    def test_whole_words_only(self):
        self.assertEqual(a.find_skills("JavaScript"), ["JavaScript"])      # not Java
        self.assertEqual(a.find_skills("MySQL"), ["MySQL"])                # not SQL
        self.assertEqual(a.find_skills("NoSQL"), ["NoSQL"])                # not SQL

    def test_everyday_words_are_not_skills(self):
        self.assertEqual(a.find_skills("We excel at a swift pace in R&D."), [])

    def test_capitalised_names_still_match(self):
        self.assertIn("Excel", a.find_skills("Advanced Excel"))
        self.assertIn("Swift", a.find_skills("Swift programming"))

    def test_symbols_in_names(self):
        found = a.find_skills("C++, C#, Node.js, CI/CD, .NET")
        for skill in ["C++", "C#", "Node.js", "CI/CD", ".NET"]:
            self.assertIn(skill, found)


class SplitJobSkillsTests(unittest.TestCase):
    def test_no_marker_means_all_required(self):
        required, preferred = a.split_job_skills("Need Python and SQL")
        self.assertEqual((required, preferred), (["Python", "SQL"], []))

    def test_nice_to_have_section(self):
        job = "Requirements: Python\nNice to have:\nDocker, Python"
        required, preferred = a.split_job_skills(job)
        self.assertEqual(required, ["Python"])
        self.assertEqual(preferred, ["Docker"])      # Python not repeated

    def test_marker_with_nothing_before_it(self):
        required, preferred = a.split_job_skills("Preferred:\nPython, SQL")
        self.assertEqual(required, ["Python", "SQL"])
        self.assertEqual(preferred, [])


class ScoringTests(unittest.TestCase):
    def test_sample_numbers(self):
        r = a.run_analysis(a.SAMPLE_RESUME, a.SAMPLE_JOB)
        self.assertEqual(r["skill_score"], 75.0)
        self.assertEqual(r["missing_skills"], ["Statistics", "Data Visualization"])
        expected = (a.SKILL_WEIGHT * r["skill_score"]
                    + a.TEXT_WEIGHT * r["text_score"])
        self.assertAlmostEqual(r["final_score"], expected)

    def test_required_skills_count_double(self):
        job = "Python\nNice to have:\nDocker"
        r = a.run_analysis("Python", job)           # has the required one
        self.assertAlmostEqual(r["skill_score"], 2 / 3 * 100)
        r = a.run_analysis("Docker", job)           # has only the nice-to-have
        self.assertAlmostEqual(r["skill_score"], 1 / 3 * 100)

    def test_missing_skills_most_important_first(self):
        r = a.run_analysis("nothing", "Python\nNice to have:\nDocker, SQL")
        self.assertEqual(r["missing_skills"][0], "Python")

    def test_job_without_known_skills(self):
        r = a.run_analysis("Python", "Head chef: cooking and menus")
        self.assertEqual(r["required_skills"], [])

    def test_status_boundaries(self):
        self.assertEqual(a.get_status(80)[0], "Excellent Match")
        self.assertEqual(a.get_status(79.9)[0], "Good Match")
        self.assertEqual(a.get_status(60)[0], "Good Match")
        self.assertEqual(a.get_status(59.9)[0], "Needs Improvement")

    def test_text_similarity_edges(self):
        self.assertGreater(a.calculate_text_similarity("python sql", "python sql"), 99)
        self.assertEqual(a.calculate_text_similarity("apple pie", "quantum physics"), 0)
        self.assertEqual(a.calculate_text_similarity("", ""), 0.0)       # no crash
        self.assertEqual(a.calculate_text_similarity("123", "456"), 0.0)

    def test_invisible_characters_do_not_hide_skills(self):
        resume = "Power\u00a0BI and Py\u200bthon"          # nbsp + zero-width space
        self.assertEqual(set(a.run_analysis(resume, "Power BI, Python")
                             ["matched_skills"]), {"Power BI", "Python"})

    def test_ligatures_are_normalised(self):
        self.assertEqual(a.clean_text("\ufb01nd \ufb02ow"), "find flow")


class ExplanationTests(unittest.TestCase):
    def test_near_perfect(self):
        self.assertIn("every skill", a.build_explanation(100, 98, [], []))

    def test_skills_are_the_problem(self):
        text = a.build_explanation(40, 90, ["SQL", "Docker"], [])
        self.assertIn("missing skills", text)
        self.assertIn("SQL", text)

    def test_wording_is_the_problem(self):
        self.assertIn("wording", a.build_explanation(100, 30, [], []))


class ResumeCheckTests(unittest.TestCase):
    def test_good_resume_has_no_tips(self):
        self.assertEqual(a.check_resume(GOOD_RESUME), [])

    def test_sample_resume_tips(self):
        tips = " ".join(a.check_resume(a.SAMPLE_RESUME))
        self.assertIn("email", tips)
        self.assertIn("numbers", tips)
        self.assertIn("short", tips)

    def test_weak_openings(self):
        tips = a.check_resume(GOOD_RESUME + "\n- Responsible for reports\n")
        self.assertTrue(any("Weak openings" in t for t in tips))

    def test_missing_section(self):
        tips = a.check_resume(GOOD_RESUME.replace("Education", "Hobbies"))
        self.assertTrue(any("'Education'" in t for t in tips))


class LearningPlanTests(unittest.TestCase):
    def test_plan_order_tags_and_hours(self):
        lines, hours = a.build_learning_plan(["SQL", "Docker"], ["Docker"])
        self.assertTrue(lines[0].startswith("1. SQL (required)"))
        self.assertTrue(lines[1].startswith("2. Docker (nice-to-have)"))
        self.assertEqual(hours, 25 + 10)

    def test_soft_skill_has_no_hours(self):
        lines, hours = a.build_learning_plan(["Communication"], [])
        self.assertEqual(hours, 0)
        self.assertIn("No course needed", lines[0])

    def test_links(self):
        self.assertEqual(a.get_learning_links("Communication"), [])
        links = a.get_learning_links("C++")
        self.assertTrue(any("C%2B%2B" in url for _, url in links))
        self.assertTrue(links[0][1].startswith("https://"))


class ReadFileTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def path(self, name):
        return os.path.join(self.tmp.name, name)

    def write(self, name, data):
        with open(self.path(name), "wb") as f:
            f.write(data)
        return self.path(name)

    def test_txt_encodings(self):
        text = "Résumé: Python"
        self.assertEqual(a.read_file(self.write("a.txt", text.encode("utf-8"))), text)
        self.assertEqual(a.read_file(self.write("b.txt", b"\xef\xbb\xbf" + text.encode("utf-8"))), text)
        self.assertEqual(a.read_file(self.write("c.txt", text.encode("utf-16"))), text)
        self.assertEqual(a.read_file(self.write("d.txt", text.encode("cp1252"))), text)

    def test_binary_renamed_to_txt(self):
        with self.assertRaises(a.FileReadError):
            a.read_file(self.write("junk.txt", bytes(range(0, 32)) * 50))

    def test_unsupported_types(self):
        with self.assertRaises(a.FileReadError):
            a.read_file(self.write("x.png", b"123"))
        with self.assertRaises(a.FileReadError) as ctx:
            a.read_file(self.write("old.doc", b"123"))
        self.assertIn(".docx", str(ctx.exception))

    def test_too_big(self):
        path = self.write("big.txt", b"python " * 100)
        with mock.patch.object(a, "MAX_FILE_MB", 0):
            with self.assertRaises(a.FileReadError):
                a.read_file(path)

    def test_docx_reads_paragraphs_and_tables(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Hello resume")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Python"
        table.cell(0, 1).text = "SQL"
        doc.save(self.path("r.docx"))
        text = a.read_file(self.path("r.docx"))
        self.assertIn("Hello resume", text)
        self.assertIn("Python", text)
        self.assertIn("SQL", text)

    def test_damaged_docx(self):
        with self.assertRaises(a.FileReadError):
            a.read_file(self.write("bad.docx", b"this is not a zip file"))

    def test_damaged_pdf(self):
        with self.assertRaises(a.FileReadError):
            a.read_file(self.write("bad.pdf", b"%PDF-1.4 garbage"))

    @unittest.skipUnless(a.HAS_REPORTLAB, "reportlab not installed")
    def test_pdf_with_text(self):
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(self.path("t.pdf"))
        c.drawString(72, 750, "Python and SQL experience for a data analyst role")
        c.save()
        self.assertIn("Python", a.read_file(self.path("t.pdf")))

    @unittest.skipUnless(a.HAS_REPORTLAB, "reportlab not installed")
    def test_blank_pdf_without_ocr(self):
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(self.path("blank.pdf"))
        c.showPage()
        c.save()
        with mock.patch.object(a, "ocr_available", return_value=False), \
             mock.patch.object(a, "_ocr_problem", "packages"):
            with self.assertRaises(a.FileReadError) as ctx:
                a.read_file(self.path("blank.pdf"))
        self.assertIn("scan", str(ctx.exception))
        self.assertIn("pip install pytesseract", str(ctx.exception))

    def test_help_message_names_the_missing_part(self):
        with mock.patch.object(a, "_ocr_problem", "program"):
            self.assertIn("Tesseract program", a.ocr_help_message())
            self.assertNotIn("pip install", a.ocr_help_message())
        with mock.patch.object(a, "_ocr_problem", "packages"):
            self.assertIn("pip install", a.ocr_help_message())

    @unittest.skipUnless(a.HAS_REPORTLAB, "reportlab not installed")
    def test_blank_pdf_uses_ocr_when_available(self):
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(self.path("scan.pdf"))
        c.showPage()
        c.save()
        with mock.patch.object(a, "ocr_available", return_value=True), \
             mock.patch.object(a, "ocr_pdf", return_value="Python resume text"):
            self.assertEqual(a.read_file(self.path("scan.pdf")), "Python resume text")


class PathTests(unittest.TestCase):
    def test_running_from_source_uses_the_code_folder(self):
        with mock.patch.object(a.sys, "frozen", False, create=True):
            self.assertEqual(a.find_data_folder(), a.find_resource_folder())

    def test_exe_keeps_history_in_appdata(self):
        with tempfile.TemporaryDirectory() as tmp, \
             mock.patch.object(a.sys, "frozen", True, create=True), \
             mock.patch.dict(os.environ, {"APPDATA": tmp}):
            folder = a.find_data_folder()
            self.assertEqual(folder, os.path.join(tmp, "ResumeJobAnalyzer"))
            self.assertTrue(os.path.isdir(folder))

    def test_exe_reads_skills_from_the_bundle(self):
        with mock.patch.object(a.sys, "_MEIPASS", "/bundle", create=True):
            self.assertEqual(a.find_resource_folder(), "/bundle")

    def test_data_files_exist_next_to_the_code(self):
        self.assertTrue(os.path.isfile(a.SKILLS_FILE))
        self.assertTrue(os.path.isfile(a.RESOURCES_FILE))


class OcrSetupTests(unittest.TestCase):
    def setUp(self):
        try:
            import pytesseract
        except ImportError:
            self.skipTest("pytesseract not installed")
        self.pytesseract = pytesseract

    def test_finds_tesseract_in_the_default_windows_folder(self):
        fake = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        env = {k: v for k, v in os.environ.items() if k != "TESSERACT_CMD"}
        with mock.patch.object(a, "_ocr_ready", None), \
             mock.patch.object(a.shutil, "which", return_value=None), \
             mock.patch.dict(os.environ, env, clear=True), \
             mock.patch.object(a.os.path, "isfile", side_effect=lambda p: p == fake), \
             mock.patch.object(self.pytesseract, "get_tesseract_version", return_value="5.0"), \
             mock.patch.object(self.pytesseract.pytesseract, "tesseract_cmd", "tesseract"):
            self.assertTrue(a.ocr_available())
            self.assertEqual(self.pytesseract.pytesseract.tesseract_cmd, fake)

    def test_reports_program_missing(self):
        env = {k: v for k, v in os.environ.items() if k != "TESSERACT_CMD"}
        with mock.patch.object(a, "_ocr_ready", None), \
             mock.patch.object(a, "_ocr_problem", ""), \
             mock.patch.object(a.shutil, "which", return_value=None), \
             mock.patch.dict(os.environ, env, clear=True), \
             mock.patch.object(a.os.path, "isfile", return_value=False), \
             mock.patch.object(self.pytesseract, "get_tesseract_version",
                               side_effect=EnvironmentError("not found")):
            self.assertFalse(a.ocr_available())
            self.assertEqual(a._ocr_problem, "program")


class HistoryTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.file = os.path.join(self.tmp.name, "history.json")

    def entry(self, signature="s1", job_key="j1", score=50.0, matched=None):
        return {"signature": signature, "job_key": job_key, "time": "2026-01-01 10:00",
                "final_score": score, "skill_score": score, "text_score": score,
                "matched_skills": matched or [], "missing_skills": []}

    def test_save_and_load_round_trip(self):
        entries = [self.entry("a"), self.entry("b")]
        a.save_history(entries, self.file)
        self.assertEqual(a.load_history(self.file), entries)

    def test_missing_file_is_empty_history(self):
        self.assertEqual(a.load_history(self.file), [])

    def test_damaged_file_is_set_aside(self):
        with open(self.file, "w") as f:
            f.write("{not json")
        self.assertEqual(a.load_history(self.file), [])
        self.assertTrue(os.path.exists(self.file + ".corrupt"))

    def test_wrong_shape_is_set_aside(self):
        with open(self.file, "w") as f:
            json.dump({"hello": 1}, f)
        self.assertEqual(a.load_history(self.file), [])

    def test_add_entry_saves_and_trims(self):
        with mock.patch.object(a, "HISTORY_FILE", self.file), \
             mock.patch.object(a, "MAX_HISTORY", 3), \
             mock.patch.object(a, "analysis_history", []):
            for i in range(5):
                self.assertIsNone(a.add_history_entry(self.entry(f"s{i}")))
            self.assertEqual([e["signature"] for e in a.analysis_history], ["s2", "s3", "s4"])
            self.assertEqual(len(a.load_history(self.file)), 3)

    def test_save_failure_is_reported_not_raised(self):
        with mock.patch.object(a, "HISTORY_FILE", os.path.join(self.tmp.name, "nope", "h.json")), \
             mock.patch.object(a, "analysis_history", []):
            self.assertIsInstance(a.add_history_entry(self.entry()), OSError)

    def test_find_previous_run(self):
        entries = [self.entry("s1", "jobA"), self.entry("s2", "jobB"), self.entry("s3", "jobA")]
        self.assertEqual(a.find_previous_run(entries, "jobA", "s9")["signature"], "s3")
        self.assertEqual(a.find_previous_run(entries, "jobA", "s3")["signature"], "s1")
        self.assertIsNone(a.find_previous_run(entries, "jobC", "s9"))
        self.assertIsNone(a.find_previous_run(entries, None, "s9"))

    def test_describe_change(self):
        previous = self.entry(score=58.0, matched=["Python"])
        text = a.describe_change(71.4, ["Python", "SQL"], previous)
        self.assertIn("58.0% -> 71.4%", text)
        self.assertIn("+13.4", text)
        self.assertIn("gained: SQL", text)
        self.assertIn("-10.0", a.describe_change(48.0, [], previous))

    def test_describe_change_when_nothing_changed(self):
        previous = self.entry(score=68.54, matched=["Python"])
        for score in (68.54, 68.52, 68.57):        # tiny differences, both signs
            text = a.describe_change(score, ["Python"], previous)
            self.assertIn("No change", text)
            self.assertNotIn("0.0 points", text)
        self.assertIn("+0.2", a.describe_change(68.74, ["Python"], previous))

    def test_job_key_ignores_spacing_and_case(self):
        self.assertEqual(a.make_job_key("Python  SQL\n"), a.make_job_key("python sql"))
        self.assertNotEqual(a.make_job_key("python"), a.make_job_key("java"))


class RankJobsTests(unittest.TestCase):
    def test_sorted_best_first_and_unrecognised_last(self):
        jobs = [
            {"name": "Chef", "text": "cooking and menus"},
            {"name": "ML", "text": "Python, Machine Learning, TensorFlow, Docker, AWS"},
            {"name": "Analyst", "text": a.SAMPLE_JOB},
        ]
        ranked = a.rank_jobs(a.SAMPLE_RESUME, jobs)
        self.assertEqual([r["name"] for r in ranked], ["Analyst", "ML", "Chef"])
        self.assertFalse(ranked[-1]["recognized"])
        self.assertEqual(ranked[0]["job_index"], 2)

    def test_job_title(self):
        self.assertEqual(a.job_title_from_text("\n\n  Data Analyst \nmore"), "Data Analyst")
        self.assertEqual(a.job_title_from_text("   "), "Untitled job")
        self.assertTrue(a.job_title_from_text("x" * 100).endswith("..."))


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.result = a.run_analysis(a.SAMPLE_RESUME, a.SAMPLE_JOB)
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def test_text_report_sections(self):
        text = a.build_report(self.result)
        for heading in ["OVERALL MATCH", "WHY THIS SCORE", "LEARNING PLAN",
                        "LEARNING LINKS", "RESUME TIPS"]:
            self.assertIn(heading, text)
        self.assertIn("Statistics", text)

    def test_docx_report(self):
        from docx import Document
        path = os.path.join(self.tmp.name, "r.docx")
        a.write_docx_report(a.build_report_sections(self.result), path)
        body = "\n".join(p.text for p in Document(path).paragraphs)
        self.assertIn("Overall Match Score", body)

    @unittest.skipUnless(a.HAS_REPORTLAB, "reportlab not installed")
    def test_pdf_report(self):
        path = os.path.join(self.tmp.name, "r.pdf")
        a.write_pdf_report(a.build_report_sections(self.result), path)
        with open(path, "rb") as f:
            self.assertEqual(f.read(4), b"%PDF")

    def test_no_missing_skills_report(self):
        result = a.run_analysis(GOOD_RESUME, a.SAMPLE_JOB)
        self.assertIn("all recognized skills", a.build_report(result))

    def test_signature_differs_for_different_text(self):
        self.assertNotEqual(a.make_signature("a", "b"), a.make_signature("a", "c"))


if __name__ == "__main__":
    unittest.main()
