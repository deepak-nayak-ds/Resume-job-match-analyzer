"""
analyzer.py - the "brain" of the Resume & Job Match Analyzer.

Reading files, finding skills, scoring, learning plans, reports and saved
history. Nothing in this file uses tkinter, so it can be tested on its own
(see tests/test_analyzer.py). The window lives in app.py.
"""

import hashlib
import json
import os
import re
import shutil
import sys
import unicodedata
from datetime import datetime
from urllib.parse import quote_plus
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# Optional: only needed to save the report as a PDF.
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

# Optional: only needed to draw the skills chart.
try:
    from matplotlib.figure import Figure
    from matplotlib.ticker import MaxNLocator
    HAS_MATPLOTLIB = True
except ImportError:
    HAS_MATPLOTLIB = False


# ============================================================
# WHERE THE FILES LIVE
# ============================================================
def find_resource_folder():
    """Folder with skills.json and resources.json. When the program is packed
    into an .exe (PyInstaller), these files are unpacked to a temporary folder
    whose path is stored in sys._MEIPASS."""
    return getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))


def find_data_folder():
    """Folder where we are allowed to WRITE (history.json).
    Next to the code when running from source. As an .exe, the unpack folder is
    deleted on exit, so use the user's AppData folder instead."""
    if not getattr(sys, "frozen", False):
        return find_resource_folder()
    base = os.environ.get("APPDATA") or os.path.expanduser("~")
    folder = os.path.join(base, "ResumeJobAnalyzer")
    try:
        os.makedirs(folder, exist_ok=True)
    except OSError:
        folder = os.path.expanduser("~")
    return folder


# ============================================================
# CONSTANTS
# ============================================================
COLOR_GREEN = "#16803C"
COLOR_ORANGE = "#C27C0E"
COLOR_RED = "#C62828"

APP_FOLDER = find_resource_folder()
SKILLS_FILE = os.path.join(APP_FOLDER, "skills.json")
RESOURCES_FILE = os.path.join(APP_FOLDER, "resources.json")   # optional
HISTORY_FILE = os.path.join(find_data_folder(), "history.json")   # saved between runs
MAX_HISTORY = 200                                              # keep the newest 200

# File reading limits
MAX_FILE_MB = 20             # refuse bigger files
SCANNED_PDF_MIN_CHARS = 30   # less text than this = probably a scanned PDF
OCR_MAX_PAGES = 5            # OCR is slow, so only read the first pages

# Where the Tesseract program usually lives when it is not on the PATH.
# (You can also set an environment variable TESSERACT_CMD to its full path.)
TESSERACT_LOCATIONS = [
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs",
                 "Tesseract-OCR", "tesseract.exe"),
    "/opt/homebrew/bin/tesseract",
    "/usr/local/bin/tesseract",
]

# Search pages on learning platforms. {q} is replaced by the skill name.
# Search links always work, even if a course is renamed or removed.
LEARNING_PLATFORMS = [
    ("Coursera", "https://www.coursera.org/search?query={q}"),
    ("YouTube", "https://www.youtube.com/results?search_query={q}"),
    ("Udemy", "https://www.udemy.com/courses/search/?q={q}"),
]

# Final score = SKILL_WEIGHT * skill match + TEXT_WEIGHT * text similarity
SKILL_WEIGHT = 0.7
TEXT_WEIGHT = 0.3
# Raw TF-IDF similarity between two documents is naturally low (a very strong
# resume scores about 40-50%). At or above this value we treat the wording
# as a "full" match, so the 30% text part isn't unfairly dragged down.
SIMILARITY_FULL_MARK = 45

REQUIRED_WEIGHT = 2    # skills in the main requirements
PREFERRED_WEIGHT = 1   # skills listed as "nice to have" / "preferred"

# A line starting with one of these words begins the "nice to have" part.
PREFERRED_MARKER = re.compile(
    r"^\s*(nice[\s-]*to[\s-]*have|preferred|good[\s-]*to[\s-]*have|bonus|desired|optional)\b.*$",
    re.IGNORECASE | re.MULTILINE,
)

# Resume check settings
SECTION_KEYWORDS = {
    "Skills": ["skills", "technologies", "tools", "competencies"],
    "Education": ["education", "academic", "qualification", "qualifications"],
    "Experience or Projects": ["experience", "projects", "project", "internship",
                               "internships", "employment", "work history"],
}
WEAK_STARTS = re.compile(
    r"^\s*[-*\u2022\u25aa\u25cf]?\s*(responsible for|worked on|worked with|"
    r"helped with|helped|assisted with|assisted in|involved in|tasked with|"
    r"duties included|was part of)\b",
    re.IGNORECASE | re.MULTILINE,
)
EMAIL_PATTERN = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")

SAMPLE_RESUME = """Deepak Nayak
Aspiring Data Analyst

Skills:
Python, SQL, Excel, Power BI, Pandas, NumPy

Projects:
E-commerce Customer Analytics
Used SQL and Power BI to analyze sales data.

Education:
B.Tech in Computer Science Engineering
"""

SAMPLE_JOB = """Data Analyst Intern

Required Skills:
Python, SQL, Excel, Power BI, Pandas,
NumPy, Statistics, Data Visualization

Responsibilities:
Analyze datasets, prepare reports,
create dashboards, and identify trends.
"""


# ============================================================
# LOGIC (nothing here uses tkinter)
# ============================================================
class FileReadError(ValueError):
    """A problem with a file that we can explain to the user in plain words."""


_ZERO_WIDTH = dict.fromkeys(map(ord, "\u200b\u200c\u200d\u2060\ufeff\u00ad"), None)


def clean_text(text):
    """Make text from PDFs, Word and web pages predictable: ligatures (fi, fl),
    non-breaking spaces and invisible characters would otherwise stop skills
    like 'Power BI' from matching."""
    return unicodedata.normalize("NFKC", text).translate(_ZERO_WIDTH)


def read_text_file(file_path):
    """Read a .txt file, trying the common encodings in turn."""
    with open(file_path, "rb") as file:
        raw = file.read()

    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):          # UTF-16 (Notepad)
        text = raw.decode("utf-16")
    else:
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = raw.decode("latin-1", errors="replace")

    # A renamed binary file decodes to junk. Refuse it instead of analyzing it.
    odd = sum(1 for c in text if ord(c) < 32 and c not in "\r\n\t")
    if text and odd / len(text) > 0.1:
        raise FileReadError("This does not look like a text file.")
    return text


def read_docx(file_path):
    try:
        document = Document(file_path)
        parts = [p.text for p in document.paragraphs]
        # Many resumes keep skills inside tables, so read those too.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
    except Exception as error:
        raise FileReadError(
            "This Word file could not be read. It may be damaged or "
            f"password-protected. ({error})") from error
    return "\n".join(parts)


_ocr_ready = None
_ocr_problem = ""      # "packages" or "program" when OCR is not available


def ocr_available():
    """True if pytesseract, Pillow and the Tesseract program are all installed.
    The answer is remembered, so restart the app after installing them."""
    global _ocr_ready, _ocr_problem
    if _ocr_ready is None:
        try:
            import pytesseract
            from PIL import Image   # noqa: F401
        except ImportError:
            _ocr_ready, _ocr_problem = False, "packages"
            return False

        # Find the Tesseract program even if it is not on the PATH
        command = os.environ.get("TESSERACT_CMD") or shutil.which("tesseract")
        if not command:
            command = next((path for path in TESSERACT_LOCATIONS
                            if path and os.path.isfile(path)), None)
        if command:
            pytesseract.pytesseract.tesseract_cmd = command

        try:
            pytesseract.get_tesseract_version()
            _ocr_ready = True
        except Exception:
            _ocr_ready, _ocr_problem = False, "program"
    return _ocr_ready


def ocr_help_message():
    """Tell the user exactly which OCR part is missing."""
    if _ocr_problem == "packages":
        return ("The Python OCR packages are not installed. Run this in a terminal:\n"
                "  pip install pytesseract pillow\n"
                "then also install the Tesseract program, and restart this app.")
    return ("The Python packages are installed, but the Tesseract program was not "
            "found. Install Tesseract (on Windows, the installer from the "
            "'UB Mannheim' Tesseract page), then restart this app.")


def ocr_pdf(reader, max_pages=OCR_MAX_PAGES):
    """Read text out of the page images of a scanned PDF."""
    import pytesseract
    parts = []
    for page in reader.pages[:max_pages]:
        try:
            for image in page.images:
                parts.append(pytesseract.image_to_string(image.image))
        except Exception:
            continue    # an image format this PDF library cannot decode
    return "\n".join(parts)


def read_pdf(file_path):
    try:
        reader = PdfReader(file_path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise FileReadError(
                "This PDF is password-protected. Remove the password and try again.")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except FileReadError:
        raise
    except Exception as error:
        raise FileReadError(
            f"This PDF could not be read. It may be damaged. ({error})") from error

    if len(text.strip()) >= SCANNED_PDF_MIN_CHARS:
        return text

    # Almost no text: probably a scan (a photo of a page). Try OCR.
    if not ocr_available():
        raise FileReadError(
            "This PDF has no selectable text, so it looks like a scan.\n\n"
            + ocr_help_message()
            + "\n\nOr paste the text into the box by hand.")
    text = ocr_pdf(reader)
    if not text.strip():
        raise FileReadError(
            "No text could be found in this PDF, even with OCR. "
            "The scan may be too blurry.")
    return text


def read_file(file_path):
    """Return the cleaned text inside a PDF, DOCX or TXT file.
    Raises FileReadError (a ValueError) with a readable message when it can't."""
    lower = file_path.lower()

    if lower.endswith(".doc"):
        raise FileReadError("Old .doc files are not supported. "
                            "Open it in Word and save it as .docx first.")
    if not lower.endswith((".pdf", ".docx", ".txt")):
        raise FileReadError("Please select a PDF, TXT, or DOCX file.")

    if os.path.getsize(file_path) > MAX_FILE_MB * 1024 * 1024:
        raise FileReadError(f"This file is larger than {MAX_FILE_MB} MB. "
                            "A resume or job description should be much smaller.")

    if lower.endswith(".pdf"):
        text = read_pdf(file_path)
    elif lower.endswith(".docx"):
        text = read_docx(file_path)
    else:
        text = read_text_file(file_path)
    return clean_text(text)


def load_skills(path):
    """Read skills.json and return a list of (name, category, compiled regex).
    Each regex matches the skill name OR any of its aliases."""
    if not os.path.exists(path):
        raise FileNotFoundError(
            f"skills.json not found.\nPut it in the same folder as app.py:\n{path}"
        )

    with open(path, "r", encoding="utf-8") as file:
        data = json.load(file)

    case_sensitive = set(data.get("_case_sensitive", []))
    skills = []

    for category, entries in data.items():
        if category.startswith("_"):      # settings, not a category
            continue
        for name, aliases in entries.items():
            terms = sorted([name] + aliases, key=len, reverse=True)
            pattern = "|".join(re.escape(term) for term in terms)
            flags = 0 if name in case_sensitive else re.IGNORECASE
            regex = re.compile(r"(?<!\w)(?:" + pattern + r")(?!\w)", flags)
            skills.append((name, category, regex))

    return skills


SKILLS = load_skills(SKILLS_FILE)
SKILL_CATEGORY = {name: category for name, category, _ in SKILLS}


def load_resources(path):
    """Read resources.json. Returns (per-skill info, per-category info).
    The file is optional: without it the plan just uses generic advice."""
    if not os.path.exists(path):
        return {}, {}
    with open(path, "r", encoding="utf-8") as file:
        data = json.load(file)
    return data.get("skills", {}), data.get("categories", {})


SKILL_RESOURCES, CATEGORY_RESOURCES = load_resources(RESOURCES_FILE)


SKILL_REGEX = {name: regex for name, _, regex in SKILLS}


def find_skill_positions(text, skill_names):
    """Return (start, end) character positions of every mention of the skills."""
    positions = []
    for name in skill_names:
        for match in SKILL_REGEX[name].finditer(text):
            positions.append((match.start(), match.end()))
    return positions


def make_job_key(job):
    """Fingerprint of a job description (ignores spacing and capitals), so the
    same job can be recognized across runs even if the resume changed."""
    normalized = " ".join(job.split()).lower()
    return hashlib.md5(normalized.encode("utf-8")).hexdigest()


def load_history(path):
    """Read saved history. A damaged file is set aside, never crashes the app."""
    if not os.path.exists(path):
        return []
    try:
        with open(path, "r", encoding="utf-8") as file:
            data = json.load(file)
        valid = isinstance(data, list) and all(
            isinstance(item, dict) and "signature" in item and "final_score" in item
            for item in data
        )
        if not valid:
            raise ValueError("unexpected history format")
        return data
    except (OSError, ValueError):
        try:
            os.replace(path, path + ".corrupt")
        except OSError:
            pass
        return []


def save_history(entries, path):
    """Write to a temporary file first, so a crash can't leave a half-written
    history file behind."""
    temp_path = path + ".tmp"
    with open(temp_path, "w", encoding="utf-8") as file:
        json.dump(entries[-MAX_HISTORY:], file, indent=2, ensure_ascii=False)
    os.replace(temp_path, path)


analysis_history = load_history(HISTORY_FILE)


def add_history_entry(entry):
    """Add an entry and save to disk. Returns an error object if saving failed."""
    analysis_history.append(entry)
    del analysis_history[:-MAX_HISTORY]
    try:
        save_history(analysis_history, HISTORY_FILE)
    except OSError as error:
        return error
    return None


def find_previous_run(entries, job_key, signature):
    """Most recent earlier run of the SAME job with a different resume."""
    if job_key is None:
        return None
    for item in reversed(entries):
        if item.get("job_key") == job_key and item["signature"] != signature:
            return item
    return None


def describe_change(final_score, matched_skills, previous):
    """e.g. '58.0% -> 71.4% (+13.4 points) since your run on 2026-10-17 12:43.'
    or 'No change (68.5%) since your run on ...' when the score did not move."""
    diff = final_score - previous["final_score"]
    old_matched = previous.get("matched_skills", [])
    gained = [s for s in matched_skills if s not in old_matched]
    lost = [s for s in old_matched if s not in matched_skills]

    if round(diff, 1) == 0:                    # avoid showing "-0.0 points"
        text = (f"No change ({final_score:.1f}%) since your run on "
                f"{previous['time']}.")
    else:
        sign = "+" if diff > 0 else ""
        text = (f"{previous['final_score']:.1f}% -> {final_score:.1f}% "
                f"({sign}{diff:.1f} points) since your run on {previous['time']}.")
    if gained:
        text += f" Skills gained: {', '.join(gained)}."
    if lost:
        text += f" Skills no longer matched: {', '.join(lost)}."
    return text


def job_title_from_text(text, limit=45):
    """Use the first non-empty line of a pasted job description as its name."""
    for line in text.splitlines():
        line = line.strip()
        if line:
            return line if len(line) <= limit else line[:limit - 3] + "..."
    return "Untitled job"


def rank_jobs(resume, jobs):
    """Analyze one resume against many jobs. jobs = [{"name", "text"}].
    Returns a list sorted best match first. Jobs with no recognized
    skills go to the bottom, because their score would mean nothing."""
    ranked = []
    for index, job in enumerate(jobs):
        result = run_analysis(resume, job["text"])
        ranked.append({
            "job_index": index,
            "name": job["name"],
            "result": result,
            "recognized": bool(result["required_skills"]),
        })
    ranked.sort(key=lambda item: (not item["recognized"],
                                  -item["result"]["final_score"]))
    return ranked


def build_category_figure(categories, final_score):
    """Stacked bar chart: matched vs missing skills per category."""
    names = list(categories.keys())
    matched = [categories[n][0] for n in names]
    missing = [categories[n][1] - categories[n][0] for n in names]
    totals = [categories[n][1] for n in names]
    rows = list(range(len(names)))

    figure = Figure(figsize=(7.2, 4.2), dpi=100)
    axes = figure.add_subplot(111)
    axes.barh(rows, matched, color=COLOR_GREEN, label="Matched")
    axes.barh(rows, missing, left=matched, color=COLOR_RED, label="Missing")
    for row, (m, t) in enumerate(zip(matched, totals)):
        axes.text(t + 0.08, row, f"{m}/{t}", va="center", fontsize=9)

    axes.set_yticks(rows)
    axes.set_yticklabels(names)
    axes.invert_yaxis()
    axes.set_xlim(0, max(totals) + 1)
    axes.xaxis.set_major_locator(MaxNLocator(integer=True))
    axes.set_xlabel("Skills the job asks for")
    axes.set_title(f"Skills by category  |  overall match {final_score:.1f}%")
    axes.legend(loc="lower right")
    figure.tight_layout()
    return figure


def find_skills(text):
    """Return the canonical names of all known skills found in the text."""
    return [name for name, _, regex in SKILLS if regex.search(text)]


def split_job_skills(job):
    """Split job skills into (required, preferred).
    Everything after a 'Nice to have / Preferred / Bonus' line counts as
    preferred, unless the skill is also mentioned earlier."""
    marker = PREFERRED_MARKER.search(job)
    if not marker:
        return find_skills(job), []

    required = find_skills(job[:marker.start()])
    preferred = [s for s in find_skills(job[marker.start():]) if s not in required]

    if not required:               # marker found but nothing before it
        return preferred, []
    return required, preferred


def format_skills(skills, preferred):
    """Join skills into text, tagging the nice-to-have ones."""
    parts = [f"{s} (nice-to-have)" if s in preferred else s for s in skills]
    return ", ".join(parts)


def calculate_text_similarity(resume_content, job_content):
    """TF-IDF cosine similarity as a percentage (0-100)."""
    try:
        vectorizer = TfidfVectorizer(
            stop_words="english",
            ngram_range=(1, 2),     # single words AND two-word phrases
            sublinear_tf=True,      # repeated words don't dominate
        )
        matrix = vectorizer.fit_transform([resume_content, job_content])
        return cosine_similarity(matrix[0:1], matrix[1:2])[0][0] * 100
    except ValueError:
        return 0.0


def get_status(match_percentage):
    """Return (status text, color) for a match percentage."""
    if match_percentage >= 80:
        return "Excellent Match", COLOR_GREEN
    if match_percentage >= 60:
        return "Good Match", COLOR_ORANGE
    return "Needs Improvement", COLOR_RED


def build_learning_plan(missing, preferred):
    """Turn the missing-skills list into numbered steps.
    Returns (list of lines, total estimated hours)."""
    lines = []
    total_hours = 0

    for number, skill in enumerate(missing, start=1):
        info = (SKILL_RESOURCES.get(skill)
                or CATEGORY_RESOURCES.get(SKILL_CATEGORY[skill])
                or {})
        tag = "nice-to-have" if skill in preferred else "required"
        hours = info.get("hours")

        if "note" in info:
            advice = info["note"]
        else:
            start = info.get("start",
                             "a beginner course, then one small project that uses it")
            advice = f"Start with: {start}."

        time_text = f" About {hours} h." if hours else ""
        lines.append(f"{number}. {skill} ({tag}).{time_text} {advice}")
        if hours:
            total_hours += hours

    return lines, total_hours


def get_learning_links(skill):
    """Return a list of (label, url) for one skill: hand-picked links from
    resources.json first, then a search on each learning platform."""
    info = SKILL_RESOURCES.get(skill, {})
    links = [(label, url) for label, url in info.get("links", [])]

    if SKILL_CATEGORY.get(skill) == "Soft Skills":
        return links          # soft skills aren't learned from a course

    query = quote_plus(f"{skill} tutorial for beginners")
    for platform, template in LEARNING_PLATFORMS:
        links.append((f"Search {platform}", template.format(q=query)))
    return links


def check_resume(resume):
    """Look at the resume itself (not the job) and return a list of tips."""
    tips = []

    # 1. Missing sections (a short line like "Skills:" or "Education")
    headings = []
    for line in resume.splitlines():
        head = line.strip().lower().split(":")[0].strip()
        if 0 < len(head) <= 35:
            headings.append(head)
    for section, keywords in SECTION_KEYWORDS.items():
        found = any(re.search(r"\b" + k + r"\b", h) for h in headings for k in keywords)
        if not found:
            tips.append(f"No '{section}' section found. Recruiters and ATS software "
                        f"look for it, so add a clear heading.")

    # 2. Contact details
    if not EMAIL_PATTERN.search(resume):
        tips.append("No email address found. Put your email near the top.")

    # 3. Numbers / measurable results (only count real sentences, 5+ words)
    sentences = [l for l in resume.splitlines() if len(l.split()) >= 5]
    if sentences and not any(re.search(r"\d", l) for l in sentences):
        tips.append("No numbers found in your descriptions. Add measurable results, "
                    "e.g. 'cut report time by 30%' or 'analyzed 50,000 orders'.")

    # 4. Weak openings
    weak = []
    for match in WEAK_STARTS.finditer(resume):
        phrase = match.group(1).lower()
        if phrase not in weak:
            weak.append(phrase)
    if weak:
        examples = ", ".join(f"'{w}'" for w in weak[:3])
        tips.append(f"Weak openings found ({examples}). Start bullets with strong "
                    f"verbs like Built, Analyzed, Automated, Designed or Improved.")

    # 5. Length
    words = len(resume.split())
    if words < 150:
        tips.append(f"Your resume is very short ({words} words). Add projects, "
                    f"tools used and results.")
    elif words > 900:
        tips.append(f"Your resume is long ({words} words). Aim for one to two pages.")

    return tips


def run_analysis(resume, job):
    """THE one analysis function. UI, history and report all use it.
    Returns a dictionary with every result."""
    resume = clean_text(resume)
    job = clean_text(job)
    resume_skills = set(find_skills(resume))
    required, preferred = split_job_skills(job)
    job_skills = required + preferred

    weight = {s: REQUIRED_WEIGHT for s in required}
    weight.update({s: PREFERRED_WEIGHT for s in preferred})

    matched = [s for s in job_skills if s in resume_skills]
    # Missing skills: most important (required) first
    missing = sorted(
        (s for s in job_skills if s not in resume_skills),
        key=lambda s: -weight[s],
    )

    total_weight = sum(weight.values())
    matched_weight = sum(weight[s] for s in matched)
    skill_score = matched_weight / total_weight * 100 if total_weight else 0.0

    raw_similarity = calculate_text_similarity(resume, job)
    text_score = min(raw_similarity / SIMILARITY_FULL_MARK * 100, 100.0)

    final_score = SKILL_WEIGHT * skill_score + TEXT_WEIGHT * text_score
    status_text, status_color = get_status(final_score)

    # Per-category breakdown: {category: [matched, total]}
    categories = {}
    for skill in job_skills:
        entry = categories.setdefault(SKILL_CATEGORY[skill], [0, 0])
        entry[1] += 1
        if skill in resume_skills:
            entry[0] += 1

    learning_plan, learning_hours = build_learning_plan(missing, preferred)

    return {
        "learning_plan": learning_plan,
        "learning_hours": learning_hours,
        "resume_tips": check_resume(resume),
        "required_skills": required,
        "preferred_skills": preferred,
        "matched_skills": matched,
        "missing_skills": missing,
        "categories": categories,
        "skill_score": skill_score,
        "raw_similarity": raw_similarity,
        "text_score": text_score,
        "final_score": final_score,
        "explanation": build_explanation(skill_score, text_score, missing, preferred),
        "status_text": status_text,
        "status_color": status_color,
    }


def build_explanation(skill_score, text_score, missing, preferred):
    """One or two sentences explaining what is holding the score back."""
    skill_loss = SKILL_WEIGHT * (100 - skill_score)   # points lost to skills
    text_loss = TEXT_WEIGHT * (100 - text_score)      # points lost to wording

    if skill_loss < 1 and text_loss < 3:
        return ("You cover every skill in the job description and your wording "
                "is very close to it.")

    if skill_loss >= text_loss:
        top = format_skills(missing[:3], preferred)
        return (f"Your score is held back mostly by missing skills "
                f"(-{skill_loss:.1f} points). Adding {top} would help most.")

    return (f"Your score is held back mostly by wording (-{text_loss:.1f} points): "
            f"your resume uses different terms from the job description. "
            f"Reuse its key terms where they truthfully describe your experience.")


def format_breakdown(result):
    return (f"Skill Match: {result['skill_score']:.1f}% "
            f"(counts {SKILL_WEIGHT:.0%})   |   "
            f"Text Similarity: {result['text_score']:.1f}% "
            f"(counts {TEXT_WEIGHT:.0%})")


def format_categories(categories):
    """e.g. 'Programming Languages 2/3 | Databases 1/1'"""
    if not categories:
        return "None"
    return " | ".join(f"{name} {m}/{t}" for name, (m, t) in categories.items())


def format_plan(result, limit=None):
    """Learning plan as text. limit=None means every item."""
    if not result["missing_skills"]:
        return "You have all recognized skills from the job description."
    lines = result["learning_plan"]
    shown = lines if limit is None else lines[:limit]
    text = "\n".join(shown)
    if limit is not None and len(lines) > limit:
        text += f"\n(+{len(lines) - limit} more in the downloaded report)"
    if result["learning_hours"]:
        text += (f"\nRough total: about {result['learning_hours']} hours "
                 f"for all {len(lines)} skills.")
    return text


def format_tips(result):
    if not result["resume_tips"]:
        return "No major issues found. Your resume structure looks solid."
    return "\n".join(f"- {tip}" for tip in result["resume_tips"])


def build_report_sections(result):
    """The report as a list of (heading, lines, use_bullets). Text, Word and
    PDF reports are all built from this one list, so they always match."""
    preferred = result["preferred_skills"]
    matched = format_skills(result["matched_skills"], preferred) or "None"
    missing = format_skills(result["missing_skills"], preferred) or "None"

    if result["missing_skills"]:
        plan_lines = list(result["learning_plan"])
        if result["learning_hours"]:
            plan_lines.append(
                f"Rough total: about {result['learning_hours']} hours "
                f"for all {len(plan_lines)} skills.")
    else:
        plan_lines = ["You have all recognized skills from the job description."]

    link_lines = []
    for skill in result["missing_skills"]:
        links = get_learning_links(skill)
        if not links:
            link_lines.append(f"{skill}: no course needed, show it in your resume.")
            continue
        link_lines.append(skill)
        link_lines.extend(f"    {label}: {url}" for label, url in links)
    if not link_lines:
        link_lines = ["None needed."]

    tips = result["resume_tips"] or [
        "No major issues found. Your resume structure looks solid."]

    return [
        ("Overall Match", [
            f"Overall Match Score: {result['final_score']:.2f}%",
            f"Status: {result['status_text']}",
            format_breakdown(result),
            f"(Raw TF-IDF text similarity: {result['raw_similarity']:.2f}%)",
        ], False),
        ("Why This Score", [result["explanation"]], False),
        ("Skills", [
            f"Skills Matched: {len(result['matched_skills'])} | "
            f"Skills Missing: {len(result['missing_skills'])}",
            f"By Category: {format_categories(result['categories'])}",
            f"Matched Skills: {matched}",
            f"Missing Skills: {missing}",
        ], False),
        ("Learning Plan (most important first)", plan_lines, False),
        ("Learning Links", link_lines, False),
        ("Resume Tips", tips, True),
    ]


def build_report(result):
    """Plain-text version of the report."""
    lines = ["=" * 40, "       AI RESUME & JOB MATCH ANALYZER", "=" * 40, ""]
    for heading, items, bullets in build_report_sections(result):
        lines += [heading.upper(), "-" * len(heading)]
        lines += [("- " + item) if bullets else item for item in items]
        lines.append("")
    lines += ["=" * 40, "          END OF REPORT", "=" * 40]
    return "\n".join(lines) + "\n"


def write_docx_report(sections, path):
    document = Document()
    document.add_heading("AI Resume & Job Match Analyzer", 0)
    document.add_paragraph(f"Report created {datetime.now():%Y-%m-%d %H:%M}")
    for heading, items, bullets in sections:
        document.add_heading(heading, level=2)
        for item in items:
            paragraph = document.add_paragraph(
                item.strip(), style="List Bullet" if bullets else None)
            if item.startswith("    "):
                paragraph.paragraph_format.left_indent = Inches(0.35)
    document.save(path)


LINK_PATTERN = re.compile(r"(https?://[^\s<]+)")


def write_pdf_report(sections, path):
    """Needs reportlab (pip install reportlab)."""
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"],
                          fontSize=10, leading=14, spaceAfter=3)

    flow = [Paragraph(escape("AI Resume & Job Match Analyzer"), styles["Title"]),
            Paragraph(f"Report created {datetime.now():%Y-%m-%d %H:%M}", body),
            Spacer(1, 8)]

    for heading, items, bullets in sections:
        flow.append(Paragraph(escape(heading), styles["Heading2"]))
        for item in items:
            indent = 18 if item.startswith("    ") else 0
            style = ParagraphStyle("line", parent=body,
                                   leftIndent=indent + (12 if bullets else 0),
                                   bulletIndent=indent)
            text = LINK_PATTERN.sub(r'<link href="\1" color="blue">\1</link>',
                                    escape(item.strip()))
            flow.append(Paragraph(text, style, bulletText="-" if bullets else None))

    SimpleDocTemplate(path, pagesize=A4, title="Resume Analysis Report",
                      leftMargin=50, rightMargin=50,
                      topMargin=50, bottomMargin=50).build(flow)


def make_signature(resume, job):
    """A short fingerprint of the resume + job text, used to detect repeats."""
    combined = resume.strip() + "\n###\n" + job.strip()
    return hashlib.md5(combined.encode("utf-8")).hexdigest()
